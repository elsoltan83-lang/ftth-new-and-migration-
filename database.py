# -*- coding: utf-8 -*-
"""
محرك قاعدة البيانات والمزامنة مع ملف الإكسيل (SQLite & Excel Engine)
يدعم الحفظ التلقائي، القراءة، الاستيراد، والتصدير
"""
import os
import sqlite3
import pandas as pd
import openpyxl
from datetime import datetime
import config

def get_connection():
    """الحصول على اتصال بقاعدة البيانات المحلية"""
    conn = sqlite3.connect(config.DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    """تهيئة جداول قاعدة البيانات وزرع البيانات الأولية من شيت الإكسيل إن وجدت"""
    conn = get_connection()
    cursor = conn.cursor()

    # جدول مشاريع الفايبر والتحويل
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS fiber_projects (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        Senior_Sector TEXT,
        Sector TEXT,
        Region TEXT,
        Type TEXT DEFAULT 'Normal',
        New_Completed_YTD INTEGER DEFAULT 0,
        New_Survey_Done_YTD INTEGER DEFAULT 0,
        New_Last_7_Days INTEGER DEFAULT 0,
        New_Yesterday INTEGER DEFAULT 0,
        New_Pending_Survey INTEGER DEFAULT 0,
        Migration_Completed_YTD INTEGER DEFAULT 0,
        Migration_Survey_Done_YTD INTEGER DEFAULT 0,
        Migration_Last_7_Days INTEGER DEFAULT 0,
        Migration_Yesterday INTEGER DEFAULT 0,
        Migration_Pending_Survey INTEGER DEFAULT 0,
        Flagged_2026 INTEGER DEFAULT 0,
        Flagged_Last_7_Days INTEGER DEFAULT 0,
        Flagged_Yesterday INTEGER DEFAULT 0,
        last_updated TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)

    # جدول المستخدمين والصلاحيات
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        full_name TEXT NOT NULL,
        role TEXT NOT NULL,
        assigned_sector TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    );
    """)

    conn.commit()

    # إذا كان جدول المشاريع فارغاً، نقوم بالاستيراد التلقائي من 'FTTH مشروع.xlsx'
    cursor.execute("SELECT COUNT(*) FROM fiber_projects;")
    count = cursor.fetchone()[0]

    if count == 0 and os.path.exists(config.DEFAULT_EXCEL_PATH):
        seed_from_excel(config.DEFAULT_EXCEL_PATH, conn)

    conn.close()

def seed_from_excel(excel_path, conn=None):
    """قراءة ملف الإكسيل الأولي وتحميله في قاعدة البيانات"""
    should_close = False
    if conn is None:
        conn = get_connection()
        should_close = True

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))

        if len(rows) > 2:
            data_rows = rows[2:] # تخطي ترويسة التقسيم والترويسة الفرعية
            records = []
            now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            for r in data_rows:
                # التحقق من أن السطر ليس فارغاً
                if not r or not any(r):
                    continue
                senior = str(r[0]).strip() if len(r) > 0 and r[0] is not None else ""
                sector = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
                region = str(r[2]).strip() if len(r) > 2 and r[2] is not None else ""
                line_type = str(r[3]).strip() if len(r) > 3 and r[3] is not None else "Normal"

                # دالة تنظيف الأرقام
                def parse_num(val):
                    try:
                        return int(float(val)) if val is not None and str(val).strip() != "" else 0
                    except:
                        return 0

                num_vals = [parse_num(r[i]) if i < len(r) else 0 for i in range(4, 17)]

                records.append((
                    senior, sector, region, line_type,
                    *num_vals,
                    now_str
                ))

            cursor = conn.cursor()
            cursor.executemany("""
                INSERT INTO fiber_projects (
                    Senior_Sector, Sector, Region, Type,
                    New_Completed_YTD, New_Survey_Done_YTD, New_Last_7_Days, New_Yesterday, New_Pending_Survey,
                    Migration_Completed_YTD, Migration_Survey_Done_YTD, Migration_Last_7_Days, Migration_Yesterday, Migration_Pending_Survey,
                    Flagged_2026, Flagged_Last_7_Days, Flagged_Yesterday,
                    last_updated
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, records)
            conn.commit()
    except Exception as e:
        print(f"Error seeding data from Excel: {e}")
    finally:
        if should_close:
            conn.close()

def load_data(sector=None, search_term=None):
    """
    تحميل البيانات كـ pandas DataFrame
    مع تطبيق تصفية الصلاحية حسب القطاع (Sector-based Access) والبحث
    """
    conn = get_connection()
    query = "SELECT * FROM fiber_projects WHERE 1=1"
    params = []

    if sector and sector != "الكل":
        query += " AND Sector = ?"
        params.append(sector)

    if search_term:
        query += " AND (Region LIKE ? OR Sector LIKE ? OR Senior_Sector LIKE ?)"
        term = f"%{search_term}%"
        params.extend([term, term, term])

    query += " ORDER BY Senior_Sector, Sector, Region"

    df = pd.read_sql_query(query, conn, params=params)
    conn.close()
    return df

def update_from_dataframe(df):
    """
    تحديث قاعدة البيانات مباشرة بناءً على جدول تم تعديله في st.data_editor
    وتحديث ملف الإكسيل التلقائي
    """
    conn = get_connection()
    cursor = conn.cursor()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for _, row in df.iterrows():
        rec_id = row.get("id")
        if pd.isna(rec_id):
            continue

        cursor.execute("""
            UPDATE fiber_projects SET
                Senior_Sector = ?,
                Sector = ?,
                Region = ?,
                Type = ?,
                New_Completed_YTD = ?,
                New_Survey_Done_YTD = ?,
                New_Last_7_Days = ?,
                New_Yesterday = ?,
                New_Pending_Survey = ?,
                Migration_Completed_YTD = ?,
                Migration_Survey_Done_YTD = ?,
                Migration_Last_7_Days = ?,
                Migration_Yesterday = ?,
                Migration_Pending_Survey = ?,
                Flagged_2026 = ?,
                Flagged_Last_7_Days = ?,
                Flagged_Yesterday = ?,
                last_updated = ?
            WHERE id = ?
        """, (
            str(row.get("Senior_Sector", "")),
            str(row.get("Sector", "")),
            str(row.get("Region", "")),
            str(row.get("Type", "Normal")),
            int(row.get("New_Completed_YTD", 0)),
            int(row.get("New_Survey_Done_YTD", 0)),
            int(row.get("New_Last_7_Days", 0)),
            int(row.get("New_Yesterday", 0)),
            int(row.get("New_Pending_Survey", 0)),
            int(row.get("Migration_Completed_YTD", 0)),
            int(row.get("Migration_Survey_Done_YTD", 0)),
            int(row.get("Migration_Last_7_Days", 0)),
            int(row.get("Migration_Yesterday", 0)),
            int(row.get("Migration_Pending_Survey", 0)),
            int(row.get("Flagged_2026", 0)),
            int(row.get("Flagged_Last_7_Days", 0)),
            int(row.get("Flagged_Yesterday", 0)),
            now_str,
            int(rec_id)
        ))

    conn.commit()
    conn.close()

    # تحديث وتصدير شيت الإكسيل التلقائي
    sync_to_excel_file()
    return True

def add_new_record(record_dict):
    """إضافة سجل/منطقة جديدة لقاعدة البيانات"""
    conn = get_connection()
    cursor = conn.cursor()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    cursor.execute("""
        INSERT INTO fiber_projects (
            Senior_Sector, Sector, Region, Type,
            New_Completed_YTD, New_Survey_Done_YTD, New_Last_7_Days, New_Yesterday, New_Pending_Survey,
            Migration_Completed_YTD, Migration_Survey_Done_YTD, Migration_Last_7_Days, Migration_Yesterday, Migration_Pending_Survey,
            Flagged_2026, Flagged_Last_7_Days, Flagged_Yesterday,
            last_updated
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        record_dict.get("Senior_Sector", ""),
        record_dict.get("Sector", ""),
        record_dict.get("Region", ""),
        record_dict.get("Type", "Normal"),
        record_dict.get("New_Completed_YTD", 0),
        record_dict.get("New_Survey_Done_YTD", 0),
        record_dict.get("New_Last_7_Days", 0),
        record_dict.get("New_Yesterday", 0),
        record_dict.get("New_Pending_Survey", 0),
        record_dict.get("Migration_Completed_YTD", 0),
        record_dict.get("Migration_Survey_Done_YTD", 0),
        record_dict.get("Migration_Last_7_Days", 0),
        record_dict.get("Migration_Yesterday", 0),
        record_dict.get("Migration_Pending_Survey", 0),
        record_dict.get("Flagged_2026", 0),
        record_dict.get("Flagged_Last_7_Days", 0),
        record_dict.get("Flagged_Yesterday", 0),
        now_str
    ))
    conn.commit()
    new_id = cursor.lastrowid
    conn.close()

    sync_to_excel_file()
    return new_id

def delete_record(record_id):
    """حذف سجل بناء على المعرف"""
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM fiber_projects WHERE id = ?", (record_id,))
    conn.commit()
    conn.close()
    sync_to_excel_file()
    return True

def sync_to_excel_file(output_path=config.DEFAULT_EXCEL_PATH):
    """
    تصدير كافة بيانات قاعدة البيانات لشيت الإكسيل بالهيكل الاحترافي
    (تحديث ملف الإكسيل التلقائي لضمان المزامنة التامة دون تدخل يدوي)
    """
    try:
        df = load_data()
        if df.empty:
            return False

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "today"
        ws.views.sheetView[0].rightToLeft = True

        # الصف الأول: العناوين المجمعة (new / migration / flagged)
        ws.merge_cells("A1:I1")
        ws["A1"] = "new (الشبكة الجديدة)"
        ws.merge_cells("J1:N1")
        ws["J1"] = "migration (تحويل النحاس إلى فايبر)"
        ws.merge_cells("O1:Q1")
        ws["O1"] = "flagged (المستهدف)"

        # الصف الثاني: أسماء الأعمدة التفصيلية
        sub_headers = [
            "Senior Sector", "Sector", "Region", "Type",
            "Completed YTD", "Survey Done YTD", "Last 7 Days", "Yesterday", "Pending Survey",
            "Completed YTD", "Survey Done YTD", "Last 7 Days", "Yesterday", "Pending Survey",
            "Flagged 2026", "Flagged Last 7 Days", "Flagged Yesterday"
        ]
        ws.append(sub_headers)

        # الصفوف اللاحقة: البيانات
        data_cols = [
            "Senior_Sector", "Sector", "Region", "Type",
            "New_Completed_YTD", "New_Survey_Done_YTD", "New_Last_7_Days", "New_Yesterday", "New_Pending_Survey",
            "Migration_Completed_YTD", "Migration_Survey_Done_YTD", "Migration_Last_7_Days", "Migration_Yesterday", "Migration_Pending_Survey",
            "Flagged_2026", "Flagged_Last_7_Days", "Flagged_Yesterday"
        ]

        for _, row in df.iterrows():
            row_data = [row.get(col, "") for col in data_cols]
            ws.append(row_data)

        # تنسيق بسيط للخلايا
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")

        for col in range(1, 18):
            cell1 = ws.cell(row=1, column=col)
            cell2 = ws.cell(row=2, column=col)
            cell1.font = header_font
            cell1.fill = header_fill
            cell1.alignment = center_align
            cell2.font = header_font
            cell2.fill = header_fill
            cell2.alignment = center_align

        wb.save(output_path)
        return True
    except Exception as e:
        print(f"Error syncing to excel: {e}")
        return False

def reset_db_from_excel(uploaded_file):
    """إعادة تهيئة وتحديث قاعدة البيانات من ملف إكسيل مرفوع حديثاً"""
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM fiber_projects;")
    conn.commit()

    # حفظ الملف المرفوع مؤقتاً وقراءته
    temp_path = "temp_uploaded.xlsx"
    with open(temp_path, "wb") as f:
        f.write(uploaded_file.getbuffer())

    seed_from_excel(temp_path, conn)
    conn.close()

    if os.path.exists(temp_path):
        os.remove(temp_path)

    sync_to_excel_file()
    return True

def get_unique_options(column_name, sector_filter=None):
    """إرجاع قائمة بالقيم الفريدة لعمود معين لدعم الـ Dropdowns"""
    conn = get_connection()
    cursor = conn.cursor()
    if sector_filter and sector_filter != "الكل":
        cursor.execute(f"SELECT DISTINCT {column_name} FROM fiber_projects WHERE Sector = ? AND {column_name} IS NOT NULL AND {column_name} != '' ORDER BY {column_name}", (sector_filter,))
    else:
        cursor.execute(f"SELECT DISTINCT {column_name} FROM fiber_projects WHERE {column_name} IS NOT NULL AND {column_name} != '' ORDER BY {column_name}")
    vals = [r[0] for r in cursor.fetchall()]
    conn.close()
    return vals
