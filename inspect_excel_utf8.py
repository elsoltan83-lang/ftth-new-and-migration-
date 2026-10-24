import sys
import openpyxl
import pandas as pd
import json

excel_file = "FTTH مشروع.xlsx"
wb = openpyxl.load_workbook(excel_file, data_only=True)
sheet = wb["today"]

# Read raw rows from openpyxl to see merged cells / header structure
raw_rows = []
for row in sheet.iter_rows(values_only=True):
    raw_rows.append([str(v) if v is not None else "" for v in row])

with open("excel_structure.json", "w", encoding="utf-8") as f:
    json.dump({
        "sheets": wb.sheetnames,
        "first_15_rows": raw_rows[:15],
        "total_rows": len(raw_rows)
    }, f, ensure_ascii=False, indent=2)

print("Saved structure to excel_structure.json successfully, total rows:", len(raw_rows))
