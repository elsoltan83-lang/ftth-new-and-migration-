import sys
import openpyxl
import pandas as pd
import json

excel_file = "FTTH مشروع.xlsx"
try:
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    print("SHEET_NAMES:", json.dumps(wb.sheetnames, ensure_ascii=False))
    
    for sheet in wb.sheetnames:
        df = pd.read_excel(excel_file, sheet_name=sheet)
        print(f"\n=== SHEET: {sheet} ===")
        print("SHAPE:", df.shape)
        print("COLUMNS:", json.dumps(list(df.columns), ensure_ascii=False))
        print("DTYPES:", df.dtypes.to_dict())
        print("SAMPLE ROWS:")
        for i, row in enumerate(df.head(5).to_dict(orient="records")):
            print(f"Row {i}: {row}")
except Exception as e:
    print("ERROR:", e)
